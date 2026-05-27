import os
import importlib.util
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter

# ── Load database ─────────────────────────────────────────────
BASE_DIR = r"D:\Cyber_Project"
_spec    = importlib.util.spec_from_file_location(
    "database",
    os.path.join(BASE_DIR, "database", "database.py")
)
db = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(db)

# ── Colors ────────────────────────────────────────────────────
C_DARK      = 'FF0A0A0F'
C_CARD      = 'FF1A1A2E'
C_HEADER_BG = 'FF16213E'
C_ROW_ALT   = 'FF0F0F1A'
C_ROW_NORM  = 'FF0A0A0F'
C_WHITE     = 'FFFFFFFF'
C_CYAN      = 'FF00F3FF'
C_RED       = 'FFFF003C'
C_GREEN     = 'FF05FFA1'
C_YELLOW    = 'FFFCEE0A'
C_PURPLE    = 'FFD300F9'
C_PINK      = 'FFFF2D78'
C_GREY      = 'FF8B949E'
C_BORDER    = 'FF2A2A4A'

ATTACK_COLORS = {
    'DDoS':       C_RED,
    'Recon':      C_YELLOW,
    'Spoofing':   C_PURPLE,
    'BruteForce': C_PINK,
    'BENIGN':     C_GREEN,
    'UNKNOWN':    C_GREY,
}

SEVERITY_MAP = {
    'DDoS':       ('CRITICAL', C_RED),
    'BruteForce': ('HIGH',     C_PINK),
    'Spoofing':   ('HIGH',     C_PINK),
    'Recon':      ('MEDIUM',   C_YELLOW),
    'BENIGN':     ('NORMAL',   C_GREEN),
    'UNKNOWN':    ('LOW',      C_GREY),
}

# ── Helpers ───────────────────────────────────────────────────
def fill(hex_color):
    return PatternFill('solid',
                       start_color=hex_color,
                       end_color=hex_color)

def font(bold=False, color=C_WHITE, size=10, italic=False):
    return Font(name='Arial', bold=bold,
                color=color, size=size, italic=italic)

def border():
    s = Side(style='thin', color=C_BORDER)
    return Border(left=s, right=s, top=s, bottom=s)

def center():
    return Alignment(horizontal='center',
                     vertical='center', wrap_text=True)

def left_align():
    return Alignment(horizontal='left',
                     vertical='center', wrap_text=True)

def style(cell, bold=False, color=C_WHITE, size=9,
          bg=C_CARD, align='center', italic=False, bdr=True):
    cell.font      = font(bold=bold, color=color,
                          size=size, italic=italic)
    cell.fill      = fill(bg)
    cell.alignment = center() if align == 'center' else left_align()
    if bdr:
        cell.border = border()

def bg_rows(ws, max_row, max_col, color=C_DARK):
    for row in ws.iter_rows(min_row=1, max_row=max_row,
                             min_col=1, max_col=max_col):
        for cell in row:
            cell.fill = fill(color)

# ══════════════════════════════════════════════════════════════
#  MAIN EXPORT FUNCTION
# ══════════════════════════════════════════════════════════════
def generate_report(exported_by='admin'):
    """
    Generates a fully styled Excel report from live database.
    Called dynamically — data is always fresh.

    Parameters:
        exported_by : username of person who clicked export

    Returns:
        filepath : path to the saved .xlsx file
    """

    # ── Fetch live data from database ─────────────────────────
    alerts  = db.get_alerts(limit=10000)
    stats   = db.get_stats()
    users   = db.get_all_users()
    syslogs = db.get_system_logs(limit=50)

    total         = stats['total']
    threats       = stats['threats']
    benign        = stats['benign']
    attack_counts = stats['attack_counts']
    now           = datetime.now()
    export_time   = now.strftime('%Y-%m-%d %H:%M:%S')
    file_date     = now.strftime('%Y-%m-%d_%H-%M-%S')

    if total == 0:
        # Return empty report if no data
        wb_empty = Workbook()
        ws = wb_empty.active
        ws['A1'] = 'No alerts in database yet.'
        out = os.path.join(BASE_DIR, "results",
                           f"iot_ids_report_{file_date}.xlsx")
        wb_empty.save(out)
        return out

    # Compute extra stats
    avg_conf   = (sum(a['confidence'] for a in alerts) /
                  total) if total > 0 else 0
    src_ips    = [a['src_ip'] for a in alerts if a['src_ip']]
    dst_ports  = [a['dst_port'] for a in alerts if a['dst_port']]
    top_port   = max(set(dst_ports), key=dst_ports.count) if dst_ports else 'N/A'
    unique_src = len(set(src_ips))

    attack_alerts = [a for a in alerts if a['attack_type'] != 'BENIGN']
    top_attacker  = ('N/A' if not attack_alerts else
                     max(set(a['src_ip'] for a in attack_alerts),
                         key=[a['src_ip'] for a in attack_alerts].count))

    ts_list   = [a['timestamp'] for a in alerts if a['timestamp']]
    date_range = (f"{min(ts_list)}  →  {max(ts_list)}"
                  if ts_list else 'N/A')

    # Sort attack counts by count desc
    sorted_attacks = sorted(attack_counts.items(),
                             key=lambda x: x[1], reverse=True)

    wb = Workbook()

    # ══════════════════════════════════════════════════════════
    #  SHEET 1 — SUMMARY
    # ══════════════════════════════════════════════════════════
    ws1 = wb.active
    ws1.title = "📊 Summary"
    ws1.sheet_view.showGridLines = False

    for col, w in [('A',3),('B',24),('C',22),
                   ('D',22),('E',22),('F',22),('G',3)]:
        ws1.column_dimensions[col].width = w

    for r in range(1, 80):
        ws1.row_dimensions[r].height = 18

    bg_rows(ws1, 80, 7)

    # Title
    ws1.merge_cells('B2:F3')
    style(ws1['B2'], bold=True, color=C_CYAN, size=18,
          bg=C_CARD, align='center')
    ws1['B2'] = '🛡  IoT Intrusion Detection System'

    ws1.merge_cells('B4:F4')
    style(ws1['B4'], color=C_GREY, size=9,
          bg=C_CARD, align='center', italic=True)
    ws1['B4'] = 'Security Alert Report  —  Auto Generated from Live Database'

    ws1.row_dimensions[5].height = 8

    # Report info rows
    def info_row(row, label, value, val_color=C_WHITE):
        ws1.merge_cells(f'B{row}:C{row}')
        style(ws1[f'B{row}'], color=C_GREY, size=9,
              bg=C_CARD, align='left')
        ws1[f'B{row}'] = label
        ws1.merge_cells(f'D{row}:F{row}')
        style(ws1[f'D{row}'], bold=True, color=val_color,
              size=9, bg=C_CARD, align='left')
        ws1[f'D{row}'] = value

    info_row(6,  '📅  Report Generated',  export_time,       C_CYAN)
    info_row(7,  '📆  Data Range',        date_range,        C_WHITE)
    info_row(8,  '👤  Exported By',
             f"{exported_by}  ({next((u['role'] for u in users if u['username']==exported_by), 'user')})",
             C_GREEN)
    info_row(9,  '🗄   Database',
             r'D:\Cyber_Project\database\ids.db',           C_WHITE)
    info_row(10, '📊  Total Records',     str(total),        C_CYAN)
    info_row(11, '🔴  Total Threats',     str(threats),      C_RED)
    info_row(12, '🟢  Benign Traffic',    str(benign),       C_GREEN)
    info_row(13, '📡  Unique Sources',    str(unique_src),   C_WHITE)
    info_row(14, '🎯  Top Targeted Port', str(top_port),     C_YELLOW)
    info_row(15, '📈  Avg Confidence',    f'{avg_conf:.1%}', C_CYAN)

    ws1.row_dimensions[16].height = 8

    # Stat cards
    def stat_card(col_start, row, label, value, color):
        c1 = get_column_letter(col_start)
        c2 = get_column_letter(col_start + 1)
        ws1.merge_cells(f'{c1}{row}:{c2}{row}')
        style(ws1[f'{c1}{row}'], color=C_GREY, size=8,
              bg=C_CARD, align='center')
        ws1[f'{c1}{row}'] = label
        ws1.merge_cells(f'{c1}{row+1}:{c2}{row+1}')
        style(ws1[f'{c1}{row+1}'], bold=True, color=color,
              size=15, bg=C_CARD, align='center')
        ws1[f'{c1}{row+1}'] = value
        ws1.row_dimensions[row+1].height = 26

    stat_card(2, 17, 'TOTAL ALERTS',    total,
              C_CYAN)
    stat_card(4, 17, 'THREATS',         threats,
              C_RED)
    stat_card(6, 17, 'BENIGN',          benign,
              C_GREEN)
    stat_card(2, 19, 'AVG CONFIDENCE',  f'{avg_conf:.1%}',
              C_YELLOW)
    stat_card(4, 19, 'UNIQUE SOURCES',  unique_src,
              C_PURPLE)
    stat_card(6, 19, 'TOP PORT',        top_port,
              C_CYAN)

    ws1.row_dimensions[21].height = 8

    # Attack breakdown table
    ws1.merge_cells('B22:F22')
    style(ws1['B22'], bold=True, color=C_CYAN, size=10,
          bg=C_HEADER_BG, align='left')
    ws1['B22'] = '  ATTACK TYPE BREAKDOWN'

    h_cols = ['B','C','D','E','F']
    for h, c in zip(['Attack Type','Count','Percentage',
                     'Avg Confidence','Severity'], h_cols):
        style(ws1[f'{c}23'], bold=True, color=C_WHITE,
              size=9, bg=C_HEADER_BG, align='center')
        ws1[f'{c}23'] = h
        ws1.row_dimensions[23].height = 20

    for i, (atype, count) in enumerate(sorted_attacks):
        r        = 24 + i
        pct      = count / total if total > 0 else 0
        a_alerts = [a for a in alerts if a['attack_type'] == atype]
        avg_c    = (sum(a['confidence'] for a in a_alerts) /
                    len(a_alerts)) if a_alerts else 0
        sev, sev_color = SEVERITY_MAP.get(atype,
                                           ('MEDIUM', C_YELLOW))
        acolor   = ATTACK_COLORS.get(atype, C_WHITE)
        bg       = C_ROW_ALT if i % 2 == 0 else C_ROW_NORM

        row_data = [atype, count,
                    f'{pct:.1%}', f'{avg_c:.1%}', sev]
        for val, c in zip(row_data, h_cols):
            style(ws1[f'{c}{r}'], color=C_WHITE,
                  size=9, bg=bg, align='center')
            ws1[f'{c}{r}'] = val
            if c == 'B':
                ws1[f'{c}{r}'].font = font(bold=True,
                                           color=acolor, size=9)
            elif c == 'F':
                ws1[f'{c}{r}'].font = font(bold=True,
                                           color=sev_color, size=9)
        ws1.row_dimensions[r].height = 16

    last_row = 24 + len(sorted_attacks)
    ws1.row_dimensions[last_row].height = 8

    # Key findings — dynamic
    findings_row = last_row + 1
    ws1.merge_cells(f'B{findings_row}:F{findings_row}')
    style(ws1[f'B{findings_row}'], bold=True, color=C_CYAN,
          size=10, bg=C_HEADER_BG, align='left')
    ws1[f'B{findings_row}'] = '  KEY FINDINGS'

    threat_rate = threats / total if total > 0 else 0
    top_attack  = sorted_attacks[0][0] if sorted_attacks else 'N/A'
    top_count   = sorted_attacks[0][1] if sorted_attacks else 0

    findings = [
        (f'⚠️  {threats} of {total} events were malicious '
         f'({threat_rate:.1%} threat rate)',             C_RED),
        (f'🔴  Most common attack: {top_attack} '
         f'({top_count} events)',                        C_YELLOW),
        (f'📊  Average detection confidence: {avg_conf:.1%}', C_CYAN),
        (f'🎯  Most targeted port: {top_port}',          C_WHITE),
        (f'🌐  {unique_src} unique source IPs detected', C_WHITE),
        (f'🛡  {benign} normal/benign events recorded',  C_GREEN),
        (f'🔝  Most active attacker IP: {top_attacker}', C_PINK),
    ]

    for j, (text, color) in enumerate(findings):
        r  = findings_row + 1 + j
        bg = C_CARD if j % 2 == 0 else C_ROW_ALT
        ws1.merge_cells(f'B{r}:F{r}')
        style(ws1[f'B{r}'], color=color, size=9,
              bg=bg, align='left')
        ws1[f'B{r}'] = text
        ws1.row_dimensions[r].height = 18

    # ══════════════════════════════════════════════════════════
    #  SHEET 2 — PIE CHART
    # ══════════════════════════════════════════════════════════
    ws2 = wb.create_sheet("📈 Charts")
    ws2.sheet_view.showGridLines = False

    for col, w in [('A',3),('B',22),('C',15),
                   ('D',15),('E',22),('F',3)]:
        ws2.column_dimensions[col].width = w

    bg_rows(ws2, 50, 6)

    ws2.merge_cells('B2:E2')
    style(ws2['B2'], bold=True, color=C_CYAN,
          size=14, bg=C_CARD, align='center')
    ws2['B2'] = '📈  Attack Distribution — Live Data'

    ws2.merge_cells('B3:E3')
    style(ws2['B3'], color=C_GREY, size=8,
          bg=C_CARD, align='center', italic=True)
    ws2['B3'] = f'Generated: {export_time}  |  Total: {total} events'

    # Table headers
    for h, c in zip(['Attack Type','Count','%'], ['B','C','D']):
        style(ws2[f'{c}5'], bold=True, color=C_WHITE,
              size=9, bg=C_HEADER_BG, align='center')
        ws2[f'{c}5'] = h
        ws2.row_dimensions[5].height = 20

    # Data — fully dynamic from database
    for i, (atype, count) in enumerate(sorted_attacks):
        r      = 6 + i
        pct    = count / total if total > 0 else 0
        acolor = ATTACK_COLORS.get(atype, C_WHITE)
        bg     = C_CARD if i % 2 == 0 else C_ROW_ALT

        ws2[f'B{r}'] = atype
        style(ws2[f'B{r}'], bold=True, color=acolor,
              size=9, bg=bg, align='center')

        ws2[f'C{r}'] = count
        style(ws2[f'C{r}'], color=C_WHITE,
              size=9, bg=bg, align='center')

        ws2[f'D{r}'] = f'{pct:.1%}'
        style(ws2[f'D{r}'], bold=True, color=acolor,
              size=9, bg=bg, align='center')

        ws2.row_dimensions[r].height = 18

    data_end = 6 + len(sorted_attacks) - 1

    # Pie chart
    pie             = PieChart()
    pie.title       = f"Attack Distribution ({total} Events)"
    pie.style       = 10
    pie.width       = 16
    pie.height      = 13

    labels = Reference(ws2, min_col=2,
                       min_row=6, max_row=data_end)
    data   = Reference(ws2, min_col=3,
                       min_row=5, max_row=data_end)
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)
    pie.dataLabels             = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showCatName = True
    pie.dataLabels.showVal     = False
    ws2.add_chart(pie, "B13")

    # ══════════════════════════════════════════════════════════
    #  SHEET 3 — ALL ALERTS
    # ══════════════════════════════════════════════════════════
    ws3 = wb.create_sheet("🚨 All Alerts")
    ws3.sheet_view.showGridLines = False

    for col, w in [('A',3),('B',20),('C',20),('D',20),
                   ('E',8),('F',10),('G',14),
                   ('H',12),('I',12),('J',3)]:
        ws3.column_dimensions[col].width = w

    bg_rows(ws3, len(alerts) + 10, 10)

    ws3.merge_cells('B2:I2')
    style(ws3['B2'], bold=True, color=C_RED,
          size=13, bg=C_CARD, align='center')
    ws3['B2'] = f'🚨  All Security Alerts  ({total} records)'

    ws3.merge_cells('B3:I3')
    style(ws3['B3'], color=C_GREY, size=8,
          bg=C_CARD, align='center', italic=True)
    ws3['B3'] = (f'Exported by: {exported_by}  |  '
                 f'Date: {export_time}  |  '
                 f'Threats: {threats}  |  Benign: {benign}')

    a_headers = ['Timestamp','Source IP','Destination IP',
                 'Port','Protocol','Attack Type',
                 'Confidence','Source']
    a_cols    = ['B','C','D','E','F','G','H','I']

    for h, c in zip(a_headers, a_cols):
        style(ws3[f'{c}5'], bold=True, color=C_WHITE,
              size=9, bg=C_HEADER_BG, align='center')
        ws3[f'{c}5'] = h
        ws3.row_dimensions[5].height = 20

    # All alert rows — dynamic
    for i, alert in enumerate(alerts):
        r      = 6 + i
        atype  = alert.get('attack_type', 'UNKNOWN')
        acolor = ATTACK_COLORS.get(atype, C_WHITE)
        bg     = C_ROW_ALT if i % 2 == 0 else C_ROW_NORM
        conf   = alert.get('confidence', 0)

        # Timestamp as plain string — fixes ##### issue
        ts = str(alert.get('timestamp', ''))

        conf_color = (C_GREEN  if conf >= 0.9 else
                      C_YELLOW if conf >= 0.8 else C_RED)

        row_vals = [
            ts,
            str(alert.get('src_ip',      '')),
            str(alert.get('dst_ip',      '')),
            str(alert.get('dst_port',    '')),
            str(alert.get('protocol',    '')),
            atype,
            f"{conf:.1%}",
            str(alert.get('source',      '')),
        ]

        for val, c in zip(row_vals, a_cols):
            style(ws3[f'{c}{r}'], color=C_WHITE,
                  size=9, bg=bg, align='center')
            ws3[f'{c}{r}'] = val
            if c == 'G':
                ws3[f'{c}{r}'].font = font(bold=True,
                                           color=acolor, size=9)
            elif c == 'H':
                ws3[f'{c}{r}'].font = font(bold=True,
                                           color=conf_color, size=9)

        ws3.row_dimensions[r].height = 16

    # ══════════════════════════════════════════════════════════
    #  SHEET 4 — USER DETAILS
    # ══════════════════════════════════════════════════════════
    ws4 = wb.create_sheet("👤 User Details")
    ws4.sheet_view.showGridLines = False

    for col, w in [('A',3),('B',25),('C',22),
                   ('D',22),('E',25),('F',3)]:
        ws4.column_dimensions[col].width = w

    bg_rows(ws4, 60, 6)

    ws4.merge_cells('B2:E2')
    style(ws4['B2'], bold=True, color=C_CYAN,
          size=13, bg=C_CARD, align='center')
    ws4['B2'] = '👤  System Users & Login Details'

    ws4.merge_cells('B3:E3')
    style(ws4['B3'], color=C_GREY, size=8,
          bg=C_CARD, align='center', italic=True)
    ws4['B3'] = f'Exported by: {exported_by}  |  {export_time}'

    ws4.row_dimensions[4].height = 8

    # Users table — dynamic from database
    for h, c in zip(['Username','Role','Account Created',
                     'Last Login'], ['B','C','D','E']):
        style(ws4[f'{c}5'], bold=True, color=C_WHITE,
              size=9, bg=C_HEADER_BG, align='center')
        ws4[f'{c}5'] = h
        ws4.row_dimensions[5].height = 20

    role_colors = {'admin': C_CYAN, 'user': C_GREEN}

    for i, u in enumerate(users):
        r      = 6 + i
        bg     = C_CARD if i % 2 == 0 else C_ROW_ALT
        ucolor = role_colors.get(u.get('role','user'), C_WHITE)
        last   = u.get('last_login') or 'Never logged in'

        for val, c in zip([u.get('username',''),
                           u.get('role','').capitalize(),
                           u.get('created_at',''),
                           last], ['B','C','D','E']):
            style(ws4[f'{c}{r}'], color=C_WHITE,
                  size=9, bg=bg, align='center')
            ws4[f'{c}{r}'] = val
            if c == 'B':
                ws4[f'{c}{r}'].font = font(bold=True,
                                           color=ucolor, size=9)

        ws4.row_dimensions[r].height = 20

    last_user_row = 6 + len(users)
    ws4.row_dimensions[last_user_row].height = 8

    # Role permissions table
    ws4.merge_cells(f'B{last_user_row+1}:E{last_user_row+1}')
    style(ws4[f'B{last_user_row+1}'], bold=True,
          color=C_CYAN, size=10,
          bg=C_HEADER_BG, align='left')
    ws4[f'B{last_user_row+1}'] = '  ROLE PERMISSIONS'

    perms = [
        ('Permission',           'admin',   'user'),
        ('View Dashboard',       '✅ Yes',  '✅ Yes'),
        ('View All Alerts',      '✅ Yes',  '✅ Yes'),
        ('Export Reports',       '✅ Yes',  '✅ Yes'),
        ('Run Attack Simulator', '✅ Yes',  '✅ Yes'),
        ('Access Admin Panel',   '✅ Yes',  '❌ No'),
        ('Manage Users',         '✅ Yes',  '❌ No'),
        ('Delete Alerts',        '✅ Yes',  '❌ No'),
        ('View System Logs',     '✅ Yes',  '❌ No'),
    ]

    for j, (perm, adm, usr) in enumerate(perms):
        r  = last_user_row + 2 + j
        bg = (C_HEADER_BG if j == 0 else
              C_CARD if j % 2 == 0 else C_ROW_ALT)
        ws4.row_dimensions[r].height = 18

        for val, c in zip([perm, adm, usr],
                           ['B','C','D']):
            style(ws4[f'{c}{r}'], color=C_WHITE,
                  size=9, bg=bg, align='center')
            ws4[f'{c}{r}'] = val
            color = C_WHITE
            if val == '✅ Yes': color = C_GREEN
            if val == '❌ No':  color = C_RED
            if j == 0:          color = C_CYAN
            ws4[f'{c}{r}'].font = font(bold=(j == 0),
                                       color=color, size=9)

    # Export info block
    exp_start = last_user_row + 2 + len(perms) + 2
    ws4.merge_cells(f'B{exp_start}:E{exp_start}')
    style(ws4[f'B{exp_start}'], bold=True,
          color=C_CYAN, size=10,
          bg=C_HEADER_BG, align='left')
    ws4[f'B{exp_start}'] = '  EXPORT INFORMATION'

    export_info = [
        ('Exported By',     exported_by),
        ('Export Time',     export_time),
        ('Total Alerts',    str(total)),
        ('Total Threats',   str(threats)),
        ('Benign Events',   str(benign)),
        ('Unique Sources',  str(unique_src)),
        ('Avg Confidence',  f'{avg_conf:.1%}'),
        ('Top Port',        str(top_port)),
        ('Dashboard URL',   'http://127.0.0.1:5000'),
        ('Database',        r'D:\Cyber_Project\database\ids.db'),
        ('System',          'IoT IDS Cyberpunk Dashboard v2.0'),
    ]

    for k, (label, value) in enumerate(export_info):
        r  = exp_start + 1 + k
        bg = C_CARD if k % 2 == 0 else C_ROW_ALT
        ws4.row_dimensions[r].height = 18

        ws4.merge_cells(f'B{r}:C{r}')
        style(ws4[f'B{r}'], color=C_GREY, size=9,
              bg=bg, align='left')
        ws4[f'B{r}'] = label

        ws4.merge_cells(f'D{r}:E{r}')
        style(ws4[f'D{r}'], bold=True, color=C_CYAN,
              size=9, bg=bg, align='left')
        ws4[f'D{r}'] = value

    # ── Save file ─────────────────────────────────────────────
    os.makedirs(os.path.join(BASE_DIR, "results"), exist_ok=True)
    filename = f"iot_ids_report_{file_date}.xlsx"
    filepath = os.path.join(BASE_DIR, "results", filename)
    wb.save(filepath)

    print(f"[EXPORT] ✅ Report saved: {filepath}")
    print(f"[EXPORT]    Total alerts : {total}")
    print(f"[EXPORT]    Threats      : {threats}")
    print(f"[EXPORT]    Exported by  : {exported_by}")

    return filepath


if __name__ == "__main__":
    path = generate_report(exported_by='admin')
    print(f"\nDone → {path}")